"""Deterministic (offline, rule-based) extraction for the supported public-style formats."""
from __future__ import annotations

import csv
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from ..schema import LineItem


def num(s) -> float | None:
    if s is None or s == "":
        return None
    if isinstance(s, (int, float)):
        return float(s)
    return float(re.sub(r"[^0-9.\-]", "", str(s)))


def parse_date(s) -> date | None:
    if s in (None, ""):
        return None
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%B %d, %Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(str(s).strip(), fmt).date()
        except ValueError:
            pass
    return None


def pdf_lines(path: Path) -> list[tuple[int, int, str]]:
    """(page, line_no, text) for every non-empty line."""
    out = []
    for p, page in enumerate(PdfReader(str(path)).pages, start=1):
        for i, text in enumerate((page.extract_text() or "").splitlines(), start=1):
            if text.strip():
                out.append((p, i, text.strip()))
    return out


def extract_io(path: Path) -> list[LineItem]:
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() for h in rows[0]]
    col = lambda *names: next(i for i, h in enumerate(hdr) if h in names)
    c = {k: col(*v) for k, v in {
        "io": ("io number", "io #"), "vendor": ("vendor", "publisher"), "campaign": ("campaign",),
        "placement": ("placement", "line item"), "start": ("flight start", "start date"),
        "end": ("flight end", "end date"), "rt": ("rate type", "buy type"), "rate": ("rate",),
        "units": ("contracted units", "units"), "cost": ("contracted cost", "cost"),
        "cur": ("currency",)}.items()}
    out = []
    for i, r in enumerate(rows[1:], start=2):
        if not r[c["placement"]]:
            continue
        out.append(LineItem(f"{path.name}#R{i}", "io", str(r[c["vendor"]]), str(r[c["campaign"]]),
                            str(r[c["placement"]]), parse_date(r[c["start"]]), parse_date(r[c["end"]]),
                            str(r[c["cur"]] or "USD"), str(r[c["io"]]), str(r[c["rt"]]).upper(),
                            num(r[c["rate"]]), num(r[c["units"]]), num(r[c["cost"]]),
                            citation=f"{path.name} '{ws.title}'!R{i}"))
    return out


def extract_invoice(path: Path) -> list[LineItem]:
    lines = pdf_lines(path)
    head = {}
    for _, _, t in lines:
        m = re.match(r"(Vendor|IO #|Campaign|Invoice #):\s*(.+)", t)
        if m:
            head[m.group(1)] = m.group(2).strip()
    out = []
    for page, ln, t in lines:
        parts = [p.strip() for p in t.split("|")]
        if len(parts) != 6 or parts[0] == "Placement":
            continue
        s, e = (parse_date(x) for x in parts[1].split("-"))
        out.append(LineItem(f"{path.name}#L{ln}", "invoice", head.get("Vendor", ""), head.get("Campaign", ""),
                            parts[0], s, e, parts[5], head.get("IO #", ""), "", num(parts[3]),
                            num(parts[2]), num(parts[4]), citation=f"{path.name} p{page} L{ln}"))
    return out


# header aliases across Meta / Google Ads / DSP exports
ALIASES = {
    "campaign": ["campaign name", "campaign"],
    "placement": ["ad set name", "ad group", "placement", "line item"],
    "start": ["reporting starts", "date start", "start date", "start"],
    "end": ["reporting ends", "date end", "end date", "end"],
    "impressions": ["impressions", "impr."],
    "clicks": ["link clicks", "clicks"],
    "spend": ["amount spent (usd)", "cost", "advertiser cost (usd)", "spend (usd)", "spend"],
}


def extract_delivery(path: Path, vendor: str) -> list[LineItem]:
    if path.suffix.lower() == ".pdf":
        lines = pdf_lines(path)
        hdr_line = next(t for _, _, t in lines if t.startswith("Campaign |"))
        table = [(f"L{ln}", f"p{p} L{ln}", [x.strip() for x in t.split("|")])
                 for p, ln, t in lines if "|" in t and t != hdr_line]
        header = [h.strip().lower() for h in hdr_line.split("|")]
        preamble = [t for _, _, t in lines]
    else:
        with path.open(newline="") as f:
            raw = list(csv.reader(f))
        hi = next(i for i, r in enumerate(raw) if len(r) > 3)  # skip report title lines
        header = [h.strip().lower() for h in raw[hi]]
        table = [(f"R{i + 1}", f"R{i + 1}", r) for i, r in enumerate(raw[hi + 1:], start=hi + 1) if r]
        preamble = [" ".join(r) for r in raw[:hi]]
    idx = {k: next((header.index(a) for a in al if a in header), None) for k, al in ALIASES.items()}
    # report-level period (Google Ads style) if rows carry no dates
    period = None
    for t in preamble:
        m = re.search(r"([A-Za-z]+ \d+, \d{4}|\d{2}/\d{2}/\d{4})\s*-\s*([A-Za-z]+ \d+, \d{4}|\d{2}/\d{2}/\d{4})", t)
        if m:
            period = (parse_date(m.group(1)), parse_date(m.group(2)))
    get = lambda r, k: r[idx[k]] if idx[k] is not None else None
    out = []
    for loc, cite, r in table:
        start = parse_date(get(r, "start")) or (period[0] if period else None)
        end = parse_date(get(r, "end")) or (period[1] if period else None)
        out.append(LineItem(f"{path.name}#{loc}", "delivery", vendor, get(r, "campaign"), get(r, "placement"),
                            start, end, "USD", impressions=num(get(r, "impressions")),
                            clicks=num(get(r, "clicks")), amount=num(get(r, "spend")),
                            citation=f"{path.name} {cite}"))
    return out
